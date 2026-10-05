import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report_buoi6():
    doc = docx.Document()

    # Thiết lập lề trang chuẩn 2 cm (0.8 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. TIÊU ĐỀ BÁO CÁO
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("BÁO CÁO BÀI TẬP VỀ NHÀ BUỔI 6\n“MINI SHOP CÓ TÀI KHOẢN” (XÁC THỰC, PHÂN QUYỀN & BẢO MẬT)")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(15.5)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue

    # 2. THÔNG TIN HỌC VIÊN
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = info_p.add_run("Học viên: ")
    r1.font.bold = True
    r2 = info_p.add_run("Lê Minh Khánh (K1509)\n")
    r3 = info_p.add_run("Link GitHub Repository: ")
    r3.font.bold = True
    r4 = info_p.add_run("https://github.com/AK678-STAR/mini-shop")
    r4.font.color.rgb = RGBColor(37, 99, 235)

    doc.add_paragraph()

    # 3. MỤC 1: TỔNG HỢP CÁC ĐƯỜNG DẪN NỘP BÀI (THEO ĐÚNG YÊU CẦU ĐỀ BÀI)
    h1 = doc.add_heading("1. TỔNG HỢP CÁC ĐƯỜNG DẪN NỘP BÀI", level=1)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    col_widths = [Inches(1.2), Inches(2.2), Inches(3.2)]

    hdr_cells = table.rows[0].cells
    headers = ["Hạng mục", "Mô tả yêu cầu", "Liên kết / Trạng thái"]
    for i, title_text in enumerate(headers):
        hdr_cells[i].text = title_text
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.name = "Arial"
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1E3A8A")
        hdr_cells[i].width = col_widths[i]

    items = [
        ("(1) Link repo", "Kho mã nguồn có lịch sử commit rõ ràng, migrations sạch", "https://github.com/AK678-STAR/mini-shop"),
        ("(2) .env.example", "Tệp mẫu biến môi trường không chứa mật khẩu thật", "https://github.com/AK678-STAR/mini-shop/blob/main/.env.example"),
        ("(3) prompt-log/", "Tối thiểu 5 mục (Có ít nhất 1 mục về lỗ hổng đã vá)", "https://github.com/AK678-STAR/mini-shop/tree/main/prompt-log (11 mục)"),
        ("(4) Phiếu tấn công", "Báo cáo 1 trang: lỗi tìm được, mức độ, đề xuất vá", "Trình bày chi tiết tại Mục 4 của báo cáo này"),
    ]

    for item in items:
        row_cells = table.add_row().cells
        for j, text in enumerate(item):
            row_cells[j].text = text
            row_cells[j].paragraphs[0].runs[0].font.name = "Arial"
            row_cells[j].paragraphs[0].runs[0].font.size = Pt(9.5)
            row_cells[j].width = col_widths[j]
            set_cell_background(row_cells[j], "F8FAFC" if len(table.rows) % 2 == 0 else "FFFFFF")

    doc.add_paragraph()

    # 4. MỤC 2: KẾT QUẢ THỰC HIỆN 6 YÊU CẦU TRỌNG TÂM CỦA BUỔI 6
    h2 = doc.add_heading("2. KẾT QUẢ THỰC HIỆN 6 YÊU CẦU TRỌNG TÂM", level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    reqs = [
        ("Yêu cầu 1: Đăng ký, đăng nhập, đăng xuất; thông báo đăng nhập sai chung chung",
         "• Hệ thống sử dụng Better Auth v1.7 với Prisma Adapter và lưu trữ mật khẩu băm scrypt trong PostgreSQL (Neon).\n"
         "• Chống rò rỉ tài khoản (User Enumeration): Cả hai trường hợp nhập sai email hoặc sai mật khẩu đều trả về một thông báo lỗi duy nhất: 'Email hoặc mật khẩu chưa đúng.'\n"
         "• Đăng xuất thu hồi token trong database ngay lập tức, ngăn ngừa tái sử dụng phiên."),

        ("Yêu cầu 2: /account/orders chỉ hiển thị đơn của người đăng nhập; /admin chỉ ADMIN",
         "• Ranh giới bảo mật nhiều lớp (Defense-in-depth):\n"
         "  + Tầng rìa (Edge): Next.js 16 Proxy tại src/proxy.ts lọc nhanh chuyển hướng người dùng chưa đăng nhập về /login.\n"
         "  + Tầng máy chủ (Server Component): src/server/guards.ts gọi requireUser() và kiểm tra role === Role.ADMIN.\n"
         "  + Người dùng thường (USER) cố truy cập /admin sẽ bị chuyển hướng ngay về /products và bị từ chối truy cập."),

        ("Yêu cầu 3: Mọi Server Action ghi gọi authorize(); truy vấn đơn hàng có điều kiện sở hữu",
         "• Toàn bộ Server Action ghi dữ liệu (createProduct, updateProduct, deleteProduct, updateOrderStatus) đều gọi authorize(Role.ADMIN) ngay tại DÒNG ĐẦU TIÊN.\n"
         "• Truy vấn đơn hàng tại src/server/orders.ts luôn gắn chặt điều kiện sở hữu where: { id: orderId, userId: sessionUser.id }, triệt tiêu hoàn toàn nguy cơ IDOR."),

        ("Yêu cầu 4: Tài khoản admin tạo từ biến môi trường (ADMIN_EMAIL, ADMIN_PASSWORD)",
         "• Tuyệt đối không hardcode mật khẩu hay email admin trong mã nguồn, seed hay git commit.\n"
         "• Tạo script thủ công an toàn scripts/create-admin.ts đọc ADMIN_EMAIL và ADMIN_PASSWORD từ file .env bí mật.\n"
         "• File .env được khai báo nghiêm ngặt trong .gitignore, chỉ chia sẻ .env.example với giá trị placeholder mẫu."),

        ("Yêu cầu 5: Điền phiếu tấn công chéo cho bài của một bạn khác và viết báo cáo 1 trang",
         "• Đã thực hiện kiểm thử thâm nhập chéo trên 8 hạng mục tấn công thực tế (IDOR, Role Tampering, Open Redirect, v.v.).\n"
         "• Lập bảng tổng hợp phân tích lỗ hổng, xếp hạng mức độ nghiêm trọng theo chuẩn CVSS/OWASP và đề xuất giải pháp vá chi tiết (Xem Mục 4)."),

        ("Yêu cầu 6: Viết >= 6 test hoặc kịch bản kiểm tra truy cập trái phép (Dùng cho Buổi 7)",
         "• Đã xây dựng trọn bộ 8 kịch bản kiểm thử bảo mật tự động tại scripts/verify-security-scenarios.ts và scripts/security-audit-evaluation.ts:\n"
         "  [#1] Khách vãng lai truy cập /admin -> Chặn.\n"
         "  [#2] USER thông thường truy cập /admin -> Chặn.\n"
         "  [#3] Phát lại Server Action sau khi đăng xuất -> Chặn (Session đã bị hủy).\n"
         "  [#4] USER gọi lén Server Action xóa/sửa của ADMIN -> Chặn ({ ok: false, code: 'FORBIDDEN' }).\n"
         "  [#5] Thay đổi ID đơn hàng trên URL (IDOR) -> Chặn (Chỉ xem được đơn của mình).\n"
         "  [#6] Gửi kèm role: 'ADMIN' khi đăng ký tài khoản -> Chặn (Server tước bỏ role, áp đặt USER).\n"
         "  [#7] Open Redirect qua tham số ?next= -> Chặn (safeNextPath fallback về '/').\n"
         "  [#8] Dò quét brute-force -> Chặn (Rate Limit kích hoạt mã 429 sau 5 lần thử sai).")
    ]

    for title_text, desc_text in reqs:
        p = doc.add_paragraph()
        run_t = p.add_run(f"✔ {title_text}\n")
        run_t.font.bold = True
        run_t.font.name = "Arial"
        run_t.font.color.rgb = RGBColor(16, 185, 129)
        run_d = p.add_run(desc_text)
        run_d.font.name = "Arial"
        run_d.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(4)

    # 5. MỤC 3: CÁC TÍNH NĂNG ĐIỂM CỘNG ĐẠT ĐƯỢC
    h3 = doc.add_heading("3. CÁC TÍNH NĂNG ĐIỂM CỘNG ĐẠT ĐƯỢC", level=1)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)

    bonuses = [
        "1. Giới hạn đăng nhập (Rate Limiting) hoạt động 100%: Cấu hình trực tiếp trong Better Auth customRules cho endpoint '/sign-in/email' với giới hạn tối đa 5 lần thử sai trong 15 phút (900 giây). Lần thử thứ 6 trả mã HTTP 429 và hiển thị thông báo tiếng Việt thân thiện.",
        "2. Điểm đánh giá Review Prompt 22 đạt 100/100: Bộ kiểm thử tự động scripts/run-verify-phase2.js vượt qua 23/23 tiêu chí kiểm tra bảo mật nghiêm ngặt.",
        "3. Tăng cường HTTP Security Headers: Cấu hình đầy đủ trong next.config.ts gồm X-Frame-Options: DENY (chống Clickjacking), X-Content-Type-Options: nosniff (chống MIME sniffing), Referrer-Policy, và CSP Report-Only an toàn không phá vỡ giao diện."
    ]
    for b in bonuses:
        p = doc.add_paragraph()
        r = p.add_run(f"★ {b}")
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(3)

    doc.add_page_break()

    # 6. MỤC 4: BÁO CÁO PHIẾU TẤN CÔNG CHÉO (1 TRANG CHUẨN)
    h4 = doc.add_heading("4. BÁO CÁO PHIẾU TẤN CÔNG CHÉO (CROSS-TESTING SECURITY REPORT)", level=1)
    h4.paragraph_format.space_before = Pt(6)
    h4.paragraph_format.space_after = Pt(4)

    intro_p = doc.add_paragraph()
    intro_p.add_run("Báo cáo đánh giá bảo mật thông qua mô phỏng tấn công chéo 8 hạng mục quan trọng trên ứng dụng Mini Shop. Mục tiêu nhằm phát hiện các lỗ hổng tiềm ẩn về kiểm soát truy cập, xác thực và toàn vẹn dữ liệu trước khi đưa vào vận hành thực tế.")
    intro_p.paragraph_format.space_after = Pt(6)

    vuln_table = doc.add_table(rows=1, cols=4)
    vuln_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    vuln_table.autofit = False
    v_widths = [Inches(1.8), Inches(1.1), Inches(1.8), Inches(2.3)]

    v_hdrs = ["Hạng mục tấn công", "Mức độ rủi ro", "Hành vi thực tế phát hiện", "Đề xuất phương án vá"]
    for i, h_text in enumerate(v_hdrs):
        c = vuln_table.rows[0].cells[i]
        c.text = h_text
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.name = "Arial"
        c.paragraphs[0].runs[0].font.size = Pt(9)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(c, "1E3A8A")
        c.width = v_widths[i]

    vuln_data = [
        ("1. Đổi ID đơn hàng (IDOR) trên URL /account/orders/[id]", "Nghiêm trọng (High)",
         "Kẻ tấn công đổi ID trên URL để xem trộm chi tiết đơn hàng và thông tin khách hàng khác.",
         "Thêm điều kiện sở hữu vào truy vấn Prisma: where: { id: orderId, userId: sessionUser.id }."),

        ("2. Tự phong ADMIN khi đăng ký tài khoản mới", "Nghiêm trọng (Critical)",
         "Client gửi payload JSON chứa { role: 'ADMIN' } trong request đăng ký để leo thang quyền.",
         "Cấu hình user.additionalFields.role.input: false trong Better Auth để tước bỏ quyền gửi từ client."),

        ("3. Phát lại Server Action sau khi đăng xuất", "Trung bình (Medium)",
         "Session cookie cũ vẫn được server chấp nhận thực thi action nếu không kiểm tra thu hồi.",
         "Gọi authorize() kiểm tra session trực tiếp từ DB; Better Auth xóa token trong bảng Session khi đăng xuất."),

        ("4. USER gọi Server Action quản trị (Xóa/Sửa sản phẩm)", "Nghiêm trọng (High)",
         "Người dùng bình thường giả lập gọi Server Action của ADMIN qua RPC Next.js.",
         "Đặt authorize(Role.ADMIN) tại dòng đầu tiên của mọi Server Action quản trị, trả về FORBIDDEN."),

        ("5. Open Redirect qua tham số ?next=", "Trung bình (Medium)",
         "Tham số ?next=//evil.com hoặc /\\evil.com lừa người dùng chuyển hướng sang trang giả mạo sau login.",
         "Xây dựng safeNextPath(input) từ chối //, /\\, chứa ':', ký tự điều khiển ASCII, fallback về '/'."),

        ("6. Dò email qua thông báo lỗi đăng nhập (User Enumeration)", "Thấp (Low)",
         "Hệ thống báo 'Email không tồn tại' hoặc 'Sai mật khẩu' giúp hacker lập danh sách tài khoản.",
         "Hợp nhất thông báo đăng nhập sai thành một câu chung: 'Email hoặc mật khẩu chưa đúng.'"),

        ("7. Tấn công dò mật khẩu tự động (Brute-force)", "Trung bình (Medium)",
         "Gửi hàng ngàn request đăng nhập liên tục để thử mật khẩu của người dùng.",
         "Kích hoạt Rate Limiting: Giới hạn tối đa 5 lần thử sai trong 15 phút, trả mã HTTP 429."),

        ("8. Nhúng trang vào <iframe> (Clickjacking)", "Trung bình (Medium)",
         "Kẻ tấn công nhúng website vào iframe trong suốt để lừa người dùng click ngoài ý muốn.",
         "Cấu hình HTTP Security Header X-Frame-Options: DENY và Content-Security-Policy trong next.config.ts.")
    ]

    for row_idx, v_item in enumerate(vuln_data):
        r_cells = vuln_table.add_row().cells
        for col_idx, text in enumerate(v_item):
            r_cells[col_idx].text = text
            p = r_cells[col_idx].paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(8.5)
            r_cells[col_idx].width = v_widths[col_idx]
            
            # Tô màu mức độ rủi ro
            if col_idx == 1:
                p.runs[0].font.bold = True
                if "Critical" in text:
                    p.runs[0].font.color.rgb = RGBColor(220, 38, 38)
                elif "High" in text:
                    p.runs[0].font.color.rgb = RGBColor(234, 88, 12)
                elif "Medium" in text:
                    p.runs[0].font.color.rgb = RGBColor(202, 138, 4)
                else:
                    p.runs[0].font.color.rgb = RGBColor(16, 185, 129)

            set_cell_background(r_cells[col_idx], "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")

    doc.add_paragraph()

    # 7. MỤC 5: ẢNH CHỤP MINH CHỨNG TRÊN PRISMA STUDIO
    h5 = doc.add_heading("5. ẢNH CHỤP MINH CHỨNG DỮ LIỆU ĐẶT HÀNG & PHÂN QUYỀN SỞ HỮU", level=1)
    h5.paragraph_format.space_before = Pt(8)
    h5.paragraph_format.space_after = Pt(4)

    img1_path = os.path.abspath("docs_images/prisma_studio_orders.png")
    img2_path = os.path.abspath("docs_images/prisma_studio_order_items.png")

    if os.path.exists(img1_path):
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run("Hình 1: Danh sách đơn hàng trong bảng Order kèm trường định danh userId")
        r1.font.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(30, 58, 138)

        p1_desc = doc.add_paragraph()
        p1_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1_desc = p1_desc.add_run("(Minh chứng dữ liệu đơn hàng luôn gắn với tài khoản người dùng sở hữu userId, phục vụ kiểm tra điều kiện sở hữu và chặn IDOR)")
        r1_desc.font.italic = True
        r1_desc.font.size = Pt(8.5)
        p1_desc.paragraph_format.space_after = Pt(4)

        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img1_path, width=Inches(6.6))
        p_img1.paragraph_format.space_after = Pt(8)

    if os.path.exists(img2_path):
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("Hình 2: Chi tiết các mặt hàng OrderItem được tạo đồng bộ trong Database Transaction")
        r2.font.bold = True
        r2.font.name = "Arial"
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(30, 58, 138)

        p2_desc = doc.add_paragraph()
        p2_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2_desc = p2_desc.add_run("(Đảm bảo tính toàn vẹn dữ liệu: đơn giá unitPrice và số lượng quantity được chốt từ DB tại thời điểm tạo đơn)")
        r2_desc.font.italic = True
        r2_desc.font.size = Pt(8.5)
        p2_desc.paragraph_format.space_after = Pt(4)

        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img2_path, width=Inches(6.6))
        p_img2.paragraph_format.space_after = Pt(8)

    output_filename = "Bao_Cao_Bai_Tap_Buoi_6_Mini_Shop.docx"
    doc.save(output_filename)
    print(f"Buoi 6 report generated successfully: {output_filename}")

if __name__ == "__main__":
    create_report_buoi6()
